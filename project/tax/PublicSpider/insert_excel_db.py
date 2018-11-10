#!/usr/bin/env python
# encoding: utf-8
'''
@author: liuyd
@file: insert_excel_db.py
@time: 2018/8/9 16:48
@desc:
'''
#import docx
import traceback
import xlrd
import pandas as pd
import os
import platform
from docx import Document
from tax.util import log
from tax.db.basic_db import DbHandle
from tax.db.dbconfig import table_map
from tax.city.common.common_db import generate_db_sql
from tax.model.db_config import DBSession
from tax.PublicSpider.get_content_static import HandDb

session=DBSession()
pl=platform.system()

db=DbHandle()

#对文件进行重命名
def rename_file():
    if pl=="Windows":
        filedir="D:\\data\\tax\\"
    else:
        filedir="/home/biuser/data/jiangxi_xzcf"
    exclfiles=os.listdir(filedir)
    log.crawler.info("excel file num is:%d"%len(exclfiles))
    for i in range(len(exclfiles)):
        new_name="qs_"+str(i)+".xls"
        try:
            log.crawler.info("start read file is:%s"%exclfiles[i])
            cfile=os.path.join(filedir,exclfiles[i])
            os.rename(cfile,os.path.join(filedir,new_name))
        except:
            continue

def read_excel(sheetnum,dbtable,data_source,excelfile,first=True):
    """
    excel 文件不变
    num为表的序号
    :param num:
    :return:
    """
    fheader=False
    try:
        data = xlrd.open_workbook(excelfile)
        table = data.sheets()[sheetnum]
        nrows = table.nrows
        for i in range(nrows):
            fields = table.row_values(i)
            fields=[value.strip() for value in fields if isinstance(value,str)]
            if dbtable=="tb_xzcf":
                if '纳税人名称' in fields or '行政处罚决定书文号' in fields:
                    cols = fields
                    if first:
                        for kvalue in fields:
                            if kvalue not in table_map[dbtable]:
                                print(kvalue)
                            break

                    else:
                        sheetNames = list(data.sheet_names())
                        num_sheet = len(sheetNames)
                        srow = i
                        # write_cols(cfile,num_sheet,srow)
                        num=read_data_pandas(dbtable,excelfile, srow, num_sheet, data_source)
                        break

            elif '纳税人名称' in fields or '法定代表人姓名' in fields or '纳税人名称(姓名)' in fields:
                fheader=True
                cols=fields
                #read_data_xlrd(table,dbtable,cols,i+1,nrows)
                if first:
                    for kvalue in fields:
                        if kvalue not in table_map[dbtable]:
                            print(kvalue)
                        break
                else:
                    sheetNames = list(data.sheet_names())
                    num_sheet = len(sheetNames)
                    srow = i
                    # write_cols(cfile,num_sheet,srow)
                    num=read_data_pandas(dbtable,excelfile, srow, num_sheet, data_source)
                    break
            else:
                pass
        if not  fheader:
            print("没有发现头部信息.........")

    except Exception as err:
        print(traceback.format_exc())
        log.error.info("read eror file is:%s"%excelfile)


def read_data_xlrd(table,dbtable,cols,snum,endnum):
    """
    读取表格除了表头的数据
    :param table: 表对象
    :param snum: 开始的行号
    :param endnum: 结束的行号
    :return:
    """

    db=DbHandle()
    for i in range(snum,endnum):
        log.crawler.info("read excel num is:%d"%i)
        item={'data_source':'ShanDong'}
        datas=table.row_values(i)
        if datas:
            try:
                for col in cols:
                    if col in table_map[dbtable]:
                        item[table_map[dbtable][col]]=datas[cols.index(col)]
                sql=generate_db_sql(dbtable,item)
                db.insert_db_func(sql=sql)
            except Exception as err:
                print(type(err))


def write_cols(file,num_sheet,srow):
    """
    测试发现表格不同的col名称用于修改bug
    :param file:
    :return:
    """
    for n_s in range(0, num_sheet):
        data = pd.read_excel(file, skiprows=srow,sheet_name=n_s)
        df = pd.DataFrame(data)
        index_li = list(df.columns)
        if index_li:
            colsstr=','.join(index_li)
            log.crawler.info(colsstr)


def read_data_pandas(table,file,srow,num_sheet,data_source,test=False):
    """
    :param table: 数据入库的表名
    :param file: excel文件
    :param srow: 跳过的行数
    :param num_sheet: sheet页数
    :return:
    """
    #统计实际入库的数据量
    intotalnum=0
    db=DbHandle()
    for n_s in range(0, num_sheet):
        data = pd.read_excel(file, skiprows=srow,sheet_name=n_s)
        df = pd.DataFrame(data)
        index_li = list(df.columns)
        if index_li:

            colsstr=','.join(index_li)
            log.crawler.info(file+"\n"+colsstr)
            result_dict={}
            for index in index_li:
                if index.strip() in table_map[table]:
                    colname=table_map[table][index.strip()]
                    values=[str(x).strip() for x in list(df[index])]

                    #log.crawler.info("column name is:%s,length is:%d"%(index,len(values)))
                    result_dict[colname]=values
            if "身份证件号码" in index_li:
                card_type="身份证"
            else:
                card_type=""

            cols=list(result_dict.keys())
            values_length=len(result_dict[cols[0]])
            if test:
                num=5
            else:
                num=values_length
            for i in range(num):
                try:
                    item={"data_source":data_source}
                    if "card_type" in list(table_map[table].values()):
                        item.update(card_type=card_type)
                    if table=="tb_credit":
                        item.update(credit_level='A')
                    for col in cols:
                        item[col]=result_dict[col][i]
                    sql=generate_db_sql(table,item)
                    r=db.insert_db_func(sql=sql)
                    if r:
                        intotalnum+=1
                        log.crawler.info("insert db success table is:%s"%table)
                except Exception as e:
                    continue

        else:
           continue
    return intotalnum


def handle_doc_docx(docfile,filename,fold,table,data_source):
    """
    python 操作word的函数
    :param file:
    :return:
    """
    from win32com import client as wc
    word = wc.Dispatch('Word.Application')
    doc = word.Documents.Open(docfile)
    docxfold=os.path.join(fold,'Docx')
    if not os.path.exists(docxfold):
        os.mkdir(docxfold)
    docxfile=os.path.join(docxfold,filename)
    doc.SaveAs(docxfile, 12, False, "", True, "", False, False, False, False)
    log.crawler.info("save docx success filename is:%s"%filename)
    handle_docx(docxfile,table,data_source)
    doc.Close()
    word.Quit()

def handle_docx(file,dbtable,data_source):
    print("start ...")
    rownum=0
    hd=HandDb(dbtable)
    document = Document(file)  # 读入文件
    tables = document.tables  # 获取文件中的表格集
    if tables:
        log.crawler.info("文档中表格的数量为:%d"%len(tables))
    for i in range(len(tables)):
        table=tables[i]
        rownum=len(table.rows) #获取表格的行数
        log.crawler.info("row num is:%d"%rownum)
        columnum=len(table.columns)
        #循环表格的行数进行内容的读取
        for i in range(rownum):
            """
            找到表格正文开始的位置
            """
            texts=[]
            for j in range(columnum):
                texts.append(table.cell(i,j).text)
            if '纳税人名称' in texts:
                clos=texts
                snum=i
                clostr=','.join(clos)
                with open("docx.txt","a",encoding="utf-8") as f:
                    f.write(clostr+'\n')
                break
        for i in range(snum+1,rownum):
            values=[]
            item = {"data_source":data_source}
            try:
                for j in range(columnum):
                    value=table.cell(i, j).text
                    values.append(value)
                    if clos[j] in table_map[dbtable]:
                        item[table_map[dbtable][clos[j]]]=value
                print(values)
                if  item.get('tax_type',"中安").isdigit():
                    print("error file",file)
                sql=hd.generate_sql_dict(item)
                if sql:
                    db.insert_db_func(sql=sql)
            except Exception as err:
                print(err)
    return rownum

def start_handle(fold,table,data_source):
    if pl == "Windows":
        filedir = "D:\\data\\tax\\{fold}".format(fold=fold)
    else:
        filedir = "/home/biuser/data/data/{fold}".format(fold=fold)
    files = os.listdir(filedir)
    log.crawler.info("excel file num is:%d" % len(files))
    t=0
    s=0
    for i in range(len(files)):
        file=files[i]
        log.crawler.info("start read file num is:%d"%i)
        ffxs=os.path.splitext(file)
        #将后缀为doc的文件转化成后缀为docx的文件
        if ffxs[1]=='.doc':
            if pl=="Windows":
                docfile=os.path.join(filedir,file)
                handle_doc_docx(docfile,str(i)+'.docx',filedir,table,data_source)
            else:
                continue
        elif ffxs[1]=='.xls' or ffxs[1]=='.xlsx':
            excelfile=os.path.join(filedir,file)
            read_excel(0,table,data_source,excelfile,False)



def find_file(filedir,table,data_source):
    """

    :param filedir: 文件目录
    :return:
    """
    sumrow=0
    filedir = "D:\\data\\tax\\{fold}".format(fold=filedir)
    files=os.listdir(filedir)
    log.crawler.info("files total num is:%d"%len(files))
    for i in range(len(files)):
        try:
            log.crawler.info("start read file index is:%d"%i)
            file=files[i]
            file=os.path.join(filedir,file)
            rownum=handle_docx(file,table,data_source)
            sumrow+=rownum
        except Exception as e:
            break
    print("sumrow num is:%d"%sumrow)

if __name__=='__main__':

    fold='NingXia\\tb_credit'
    data_source='NingXia'
    table='tb_credit'
    start_handle(fold,table,data_source)
